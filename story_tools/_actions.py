"""Plain Python actions for the story-automation workflow.

No decorators, no LLM. Each function does exactly one thing and is safe to
call from either the v1 CLI or the v2 agent.
"""

import base64
import json
import os

import requests
from openpyxl import load_workbook

ORG       = os.environ["AZURE_DEVOPS_ORG"]
PROJECT   = os.environ["AZURE_DEVOPS_PROJECT"]
PAT       = os.environ["AZURE_DEVOPS_PAT"]
AREA      = os.environ["AZURE_DEVOPS_AREA_PATH"]
ITERATION = os.environ["AZURE_DEVOPS_ITERATION_PATH"]

REQUIRED_COLS = [
    "EPIC", "FEATURE", "STORY", "STORY TITLE", "DESCRIPTION",
    "ACCEPTANCE CRITERIA", "STORYPOINTS", "PRIORITY",
]

_token = base64.b64encode(f":{PAT}".encode()).decode()
HEADERS = {
    "Content-Type": "application/json-patch+json",
    "Authorization": f"Basic {_token}",
}


def load_stories_from_xlsx(path: str) -> list[dict]:
    """Read an xlsx file and return its rows as a list of dicts.

    Validates that all REQUIRED_COLS are present in the header row.
    Skips fully empty rows.
    """
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    missing = [c for c in REQUIRED_COLS if c not in headers]
    if missing:
        raise ValueError(f"Missing required columns: {missing}")
    col_idx = {h: i for i, h in enumerate(headers)}
    rows: list[dict] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        rows.append({c: row[col_idx[c]] for c in REQUIRED_COLS})
    return rows


def list_dry_run_summary(rows: list[dict]) -> dict:
    """Summarise what would be created without touching Azure DevOps."""
    unique_epics = sorted({r.get("EPIC") for r in rows if r.get("EPIC")})
    unique_features = sorted({r.get("FEATURE") for r in rows if r.get("FEATURE")})
    return {
        "epic_count": len(unique_epics),
        "feature_count": len(unique_features),
        "story_count": len(rows),
        "epics": unique_epics,
        "features": unique_features,
    }


def _create_work_item(wit_type: str, fields: dict, parent_id: int | None = None) -> int:
    """POST a new work item. wit_type is 'Epic' or 'Issue' on the Basic process."""
    url = (
        f"https://dev.azure.com/{ORG}/{PROJECT}/"
        f"_apis/wit/workitems/${wit_type}?api-version=7.0"
    )
    payload = [
        {"op": "add", "path": f"/fields/{k}", "value": v}
        for k, v in fields.items() if v not in (None, "")
    ]
    if parent_id:
        payload.append({"op": "add", "path": "/relations/-", "value": {
            "rel": "System.LinkTypes.Hierarchy-Reverse",
            "url": (
                f"https://dev.azure.com/{ORG}/{PROJECT}/"
                f"_apis/wit/workitems/{parent_id}"
            ),
        }})
    res = requests.post(url, headers=HEADERS, data=json.dumps(payload))
    if not res.ok:
        raise RuntimeError(f"Azure DevOps {res.status_code}: {res.text}")
    return res.json()["id"]


def ensure_epic_on_board(epic_name: str, cache: dict[str, int]) -> int:
    """Create the Epic if not already in `cache`. Return its work item id."""
    if epic_name in cache:
        return cache[epic_name]
    epic_id = _create_work_item("Epic", {
        "System.Title": epic_name,
        "System.AreaPath": AREA,
        "System.IterationPath": ITERATION,
    })
    cache[epic_name] = epic_id
    return epic_id


def create_story_on_board(row: dict, epic_id: int | None) -> int:
    """Create one Issue (Story) optionally parented to the given Epic id.

    Required key in `row`: STORY TITLE. All other keys are optional and may
    be omitted, None, or empty string.
    """
    if not row.get("STORY TITLE"):
        raise ValueError("row is missing required key 'STORY TITLE'")
    description = row.get("DESCRIPTION") or ""
    ac = row.get("ACCEPTANCE CRITERIA")
    if ac:
        ac_lines = [
            f"<li>{line.strip()}</li>"
            for line in str(ac).splitlines()
            if line.strip()
        ]
        description += (
            "<br><br><b>Acceptance Criteria:</b><ol>" + "".join(ac_lines) + "</ol>"
        )
    feature = row.get("FEATURE")
    fields = {
        "System.Title":                     row["STORY TITLE"],
        "System.Description":               description,
        "System.AreaPath":                  AREA,
        "System.IterationPath":             ITERATION,
        "System.Tags":                      f"feature:{feature}" if feature else None,
        "Microsoft.VSTS.Scheduling.Effort": row.get("STORYPOINTS"),
        "Microsoft.VSTS.Common.Priority":   row.get("PRIORITY"),
    }
    return _create_work_item("Issue", fields, parent_id=epic_id)
