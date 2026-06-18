import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Create workbook and sheet
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Template Ingestion Stories"

# Headers
headers = ["Story ID", "Story Title", "As a", "I want", "So that", "Scenario #", "Given", "When", "Then"]
ws.append(headers)

# Style headers
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for col in range(1, len(headers)+1):
    cell = ws.cell(row=1, column=col)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment

# Data rows (all stories and scenarios)
data = [
    ["STORY-1", "Ingest Template Header Records", "Data Engineer", "ingest template header data (def.Template) from a staging source", "each template is uniquely identified and associated with the correct tenant, version, and status", "1", "A source file contains valid template records with TemplateId, Code, Name, Version, Status, etc.", "The ingestion pipeline processes the file", "Each record is inserted into def.Template with a new or matching TemplateId"],
    ["STORY-1", "Ingest Template Header Records", "Data Engineer", "ingest template header data (def.Template) from a staging source", "each template is uniquely identified and associated with the correct tenant, version, and status", "2", "A record has the same (TenantId, Code, Version) as an existing template", "The pipeline attempts to insert the record", "The record is rejected with a duplicate key error; no update is performed (unless upsert is explicitly designed)"],
    ["STORY-1", "Ingest Template Header Records", "Data Engineer", "ingest template header data (def.Template) from a staging source", "each template is uniquely identified and associated with the correct tenant, version, and status", "3", "A record has TemplateId that already exists", "The pipeline processes the record", "The existing row is updated (or rejected based on business rule) – clearly logged"],
    ["STORY-1", "Ingest Template Header Records", "Data Engineer", "ingest template header data (def.Template) from a staging source", "each template is uniquely identified and associated with the correct tenant, version, and status", "4", "TenantId is NULL", "The pipeline ingests the record", "The template is correctly stored as a global STANDARD template (global scope)"],
    ["STORY-1", "Ingest Template Header Records", "Data Engineer", "ingest template header data (def.Template) from a staging source", "each template is uniquely identified and associated with the correct tenant, version, and status", "5", "Any required column (e.g., Name) is missing or NULL", "The pipeline validates the record", "The record is moved to an error table with a clear reason (e.g., “Missing Name”)"],

    ["STORY-2", "Ingest Template Section Hierarchy", "Data Engineer", "ingest section definitions (def.TemplateSection) and their parent‑child relationships", "a complete navigable section tree is built for each template", "1", "A source file contains sections with TemplateId, SectionCode, ParentSectionId, Level, DisplayOrder", "The pipeline processes sections after the parent template exists", "All sections are inserted, and SectionPath is materialised correctly (e.g., “4.1.1” → “4/4.1/4.1.1”)"],
    ["STORY-2", "Ingest Template Section Hierarchy", "Data Engineer", "ingest section definitions (def.TemplateSection) and their parent‑child relationships", "a complete navigable section tree is built for each template", "2", "A section refers to a ParentSectionId that does not yet exist in the same batch", "The pipeline uses a two‑pass or recursive CTE load", "The pipeline either orders records by Level or rejects orphans with a clear message"],
    ["STORY-2", "Ingest Template Section Hierarchy", "Data Engineer", "ingest section definitions (def.TemplateSection) and their parent‑child relationships", "a complete navigable section tree is built for each template", "3", "A section violates the UNIQUE (TemplateId, SectionCode) constraint", "The pipeline attempts to insert a duplicate", "The record is rejected and logged as a duplicate section code"],
    ["STORY-2", "Ingest Template Section Hierarchy", "Data Engineer", "ingest section definitions (def.TemplateSection) and their parent‑child relationships", "a complete navigable section tree is built for each template", "4", "The source provides IsRequired and FillBucket values", "The pipeline ingests these flags", "The values are stored exactly as provided; FillBucket is validated against the CHECK constraint (user_mandatory, agent_auto, agent_suggest)"],

    ["STORY-3", "Seed Control Type Lookup Table", "Data Engineer", "seed and maintain the def.ControlType reference table", "TemplateControl records can reference valid control types via foreign key", "1", "A seed script contains control type definitions (Code, DisplayName, ValueKind, HasOptions, ConfigSchema)", "The deployment pipeline runs the seed script", "All control types are inserted; duplicates are skipped or updated idempotently"],
    ["STORY-3", "Seed Control Type Lookup Table", "Data Engineer", "seed and maintain the def.ControlType reference table", "TemplateControl records can reference valid control types via foreign key", "2", "A new control type (e.g., “slider”) is added in a later release", "The data engineer runs an incremental seed update", "The new record is inserted without affecting existing types"],
    ["STORY-3", "Seed Control Type Lookup Table", "Data Engineer", "seed and maintain the def.ControlType reference table", "TemplateControl records can reference valid control types via foreign key", "3", "ValueKind is not one of (text, number, date, datetime, bool, json)", "The pipeline validates the seed data", "The record is rejected with a check constraint violation error"],

    ["STORY-4", "Ingest Template Controls (Form Fields)", "Data Engineer", "ingest form control definitions (def.TemplateControl) linked to sections and control types", "each field is correctly placed in the form and has validation rules", "1", "A source file contains controls with SectionId, ControlTypeId, FieldKey, Label, DisplayOrder", "The pipeline processes controls after sections and control types are loaded", "Each control is inserted with a valid foreign key to TemplateSection and ControlType"],
    ["STORY-4", "Ingest Template Controls (Form Fields)", "Data Engineer", "ingest form control definitions (def.TemplateControl) linked to sections and control types", "each field is correctly placed in the form and has validation rules", "2", "Two controls in the same section have the same FieldKey but different ProjectId (one NULL, one specific project)", "The pipeline attempts to insert both", "Both are accepted because the unique constraint is (SectionId, FieldKey, ProjectId) – project‑specific override is allowed"],
    ["STORY-4", "Ingest Template Controls (Form Fields)", "Data Engineer", "ingest form control definitions (def.TemplateControl) linked to sections and control types", "each field is correctly placed in the form and has validation rules", "3", "ValidationJson is not valid JSON or does not match expected schema (e.g., {min,max})", "The pipeline validates the JSON", "The record is rejected with a “invalid ValidationJson” error"],
    ["STORY-4", "Ingest Template Controls (Form Fields)", "Data Engineer", "ingest form control definitions (def.TemplateControl) linked to sections and control types", "each field is correctly placed in the form and has validation rules", "4", "ConfigJson references a table_grid with columns array but missing required keys", "The pipeline validates against ControlType’s ConfigSchema", "The record is rejected with a schema validation error"],

    ["STORY-5", "Ingest Control Options", "Data Engineer", "ingest options for controls that have HasOptions = true", "dropdowns, radio groups, and checkboxes have their choice lists populated", "1", "A source file contains options with ControlId, OptionValue, OptionLabel, DisplayOrder", "The pipeline processes options after controls exist", "Each option is inserted into def.ControlOption with a valid foreign key"],
    ["STORY-5", "Ingest Control Options", "Data Engineer", "ingest options for controls that have HasOptions = true", "dropdowns, radio groups, and checkboxes have their choice lists populated", "2", "The same ControlId, OptionValue, and ProjectId combination is ingested twice", "The pipeline attempts to insert a duplicate", "The second record is rejected (unique constraint violation)"],
    ["STORY-5", "Ingest Control Options", "Data Engineer", "ingest options for controls that have HasOptions = true", "dropdowns, radio groups, and checkboxes have their choice lists populated", "3", "An option has ProjectId = NULL (template‑wide) and another option has the same ControlId, same OptionValue but ProjectId = specific GUID (tenant override)", "The pipeline processes both", "Both are stored – the UI will prefer the project‑specific option when applicable"],
    ["STORY-5", "Ingest Control Options", "Data Engineer", "ingest options for controls that have HasOptions = true", "dropdowns, radio groups, and checkboxes have their choice lists populated", "4", "The parent control’s HasOptions is false, but options are provided", "The pipeline loads options", "The load may proceed (referential integrity only), but business logic should warn about inconsistency – logged as warning"],

    ["STORY-6", "Handle Incremental & Full Load with History", "Data Engineer", "support both full refresh and incremental (changed‑only) ingestion of template metadata", "downstream consumers always see the latest published version, while audit history is preserved", "1", "A template is marked as archived in the source system", "The incremental load detects that the status changed", "The pipeline updates Status = 'archived' and IsActive = 0 in def.Template"],
    ["STORY-6", "Handle Incremental & Full Load with History", "Data Engineer", "support both full refresh and incremental (changed‑only) ingestion of template metadata", "downstream consumers always see the latest published version, while audit history is preserved", "2", "A new version of a template (same Code, new Version) is ingested", "The pipeline processes the new record", "Both versions coexist; the latest version can be identified by max(Version) per (TenantId, Code)"],
    ["STORY-6", "Handle Incremental & Full Load with History", "Data Engineer", "support both full refresh and incremental (changed‑only) ingestion of template metadata", "downstream consumers always see the latest published version, while audit history is preserved", "3", "The source system sends a control that should be removed (soft delete)", "The pipeline receives a deletion flag", "The pipeline sets IsActive = 0 for that control, instead of physically deleting the row"],
    ["STORY-6", "Handle Incremental & Full Load with History", "Data Engineer", "support both full refresh and incremental (changed‑only) ingestion of template metadata", "downstream consumers always see the latest published version, while audit history is preserved", "4", "A full refresh loads all templates for a tenant", "The pipeline truncates and reloads the tenant’s data", "Existing active records are replaced; historical data that was not in the source is soft‑deleted (or hard deleted per business rule)"],

    ["STORY-7", "Validate Referential Integrity During Ingestion", "Data Engineer", "the ingestion pipeline enforce all foreign key and check constraints", "no orphaned or invalid data enters the database", "1", "A section references a TemplateId that does not exist in def.Template", "The pipeline attempts to insert the section", "The insert is rejected; the record is sent to a dead‑letter queue with FK violation details"],
    ["STORY-7", "Validate Referential Integrity During Ingestion", "Data Engineer", "the ingestion pipeline enforce all foreign key and check constraints", "no orphaned or invalid data enters the database", "2", "A control references a SectionId that does not exist", "The pipeline processes the control", "The control is rejected – logged with “missing section”"],
    ["STORY-7", "Validate Referential Integrity During Ingestion", "Data Engineer", "the ingestion pipeline enforce all foreign key and check constraints", "no orphaned or invalid data enters the database", "3", "A control option references a ControlId that does not exist", "The pipeline attempts to insert the option", "The option is rejected with foreign key error"],
    ["STORY-7", "Validate Referential Integrity During Ingestion", "Data Engineer", "the ingestion pipeline enforce all foreign key and check constraints", "no orphaned or invalid data enters the database", "4", "The pipeline processes a batch with all dependencies pre‑loaded (templates first, then sections, then controls, then options)", "The pipeline respects the load order", "All records are inserted successfully without FK violations"],

    ["STORY-8", "Logging, Monitoring, and Error Handling", "Data Engineer", "the ingestion process produce structured logs and metrics for each batch", "failures can be diagnosed quickly and data quality issues are visible to operations", "1", "A batch of 1,000 template records is ingested", "The pipeline completes the batch", "Logs show: number of inserted, updated, rejected, and skipped records"],
    ["STORY-8", "Logging, Monitoring, and Error Handling", "Data Engineer", "the ingestion process produce structured logs and metrics for each batch", "failures can be diagnosed quickly and data quality issues are visible to operations", "2", "A record fails due to a duplicate key", "The pipeline handles the error", "The error is logged with the record’s identifier, the violating constraint name, and the timestamp"],
    ["STORY-8", "Logging, Monitoring, and Error Handling", "Data Engineer", "the ingestion process produce structured logs and metrics for each batch", "failures can be diagnosed quickly and data quality issues are visible to operations", "3", "The pipeline is executed via an orchestrator (e.g., ADF, Airflow)", "The pipeline runs", "A final status (success / partial success / failure) is emitted to the orchestrator, and a summary metric is pushed to Azure Monitor"],
    ["STORY-8", "Logging, Monitoring, and Error Handling", "Data Engineer", "the ingestion process produce structured logs and metrics for each batch", "failures can be diagnosed quickly and data quality issues are visible to operations", "4", "A transient database timeout occurs", "The pipeline retries the operation", "The retry is logged with attempt number; after 3 failures, the batch is marked as failed and an alert is triggered"],
]

for row in data:
    ws.append(row)

# Adjust column widths
column_widths = [12, 35, 18, 45, 55, 10, 55, 45, 60]
for i, width in enumerate(column_widths, 1):
    ws.column_dimensions[chr(64+i)].width = width

# Wrap text for all cells
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=9):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

# Freeze header row
ws.freeze_panes = 'A2'

# Save file
wb.save("template_ingestion_stories.xlsx")
print("✅ Excel file created: template_ingestion_stories.xlsx")